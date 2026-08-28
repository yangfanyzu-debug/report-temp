from __future__ import annotations

import zipfile
import re
from pathlib import Path

from docx import Document


DOCX_MARKERS = {"[Content_Types].xml", "word/document.xml"}


def is_docx(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as archive:
            return DOCX_MARKERS.issubset(archive.namelist())
    except (OSError, zipfile.BadZipFile):
        return False


def extract_text_and_tables(path: Path) -> dict[str, object]:
    document = Document(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    headings: list[dict[str, object]] = []
    toc_entries: list[str] = []
    toc_field_found = False

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style_name = str(paragraph.style.name or "")
        heading_match = re.match(r"^(?:Heading|标题)\s*(\d+)$", style_name, re.IGNORECASE)
        if text and heading_match:
            headings.append(
                {"text": text, "level": int(heading_match.group(1)), "style": style_name}
            )

        field_text = " ".join(
            str(node.text or "") for node in paragraph._p.xpath(".//w:instrText")
        )
        if re.search(r"\bTOC\b", field_text, re.IGNORECASE):
            toc_field_found = True
        if text and re.match(r"^(?:TOC|目录)\s*\d+$", style_name, re.IGNORECASE):
            toc_entries.append(text)
    tables: list[list[list[str]]] = []
    for table in document.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append(rows)

    return {
        "paragraphs": paragraphs,
        "tables": tables,
        "tableCount": len(tables),
        "paragraphCount": len(paragraphs),
        "headings": headings,
        "tocEntries": toc_entries,
        "tocAvailable": bool(toc_field_found and toc_entries),
    }
