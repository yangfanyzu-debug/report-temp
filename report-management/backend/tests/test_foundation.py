from __future__ import annotations

import io
import sys
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path


BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))

from app import create_app  # noqa: E402
from app.services.docx_tools import extract_text_and_tables, is_docx  # noqa: E402
from app.services.file_naming import build_versioned_filename  # noqa: E402


def write_docx(path: Path, content: str = "报告内容") -> None:
    from docx import Document

    document = Document()
    document.add_heading("性能容量报告", level=1)
    document.add_paragraph(content)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "平均值"
    table.cell(1, 0).text = "响应时间"
    table.cell(1, 1).text = "120ms"
    document.save(path)


class FoundationTest(unittest.TestCase):
    def test_health_check(self):
        app = create_app({"TESTING": True})
        client = app.test_client()

        response = client.get("/health")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json, {"service": "report-management", "status": "ok"})

    def test_docx_detection_and_extraction(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            docx_path = Path(temporary_dir) / "报告.docx"
            write_docx(docx_path)

            self.assertTrue(is_docx(docx_path))
            extracted = extract_text_and_tables(docx_path)
            self.assertIn("性能容量报告", extracted["paragraphs"])
            self.assertEqual(extracted["tableCount"], 1)
            self.assertEqual(extracted["tables"][0][1], ["响应时间", "120ms"])

    def test_rejects_invalid_docx_zip(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            invalid_path = Path(temporary_dir) / "fake.docx"
            with zipfile.ZipFile(invalid_path, "w") as archive:
                archive.writestr("not-word.txt", "bad")

            self.assertFalse(is_docx(invalid_path))

    def test_versioned_filename_does_not_reuse_original_name(self):
        filename = build_versioned_filename(
            "credit-card-center",
            "2025年08月",
            2,
            "报告 修订版.docx",
            now=datetime(2026, 8, 11, 12, 30, 45, 123456),
        )

        self.assertEqual(
            filename,
            "credit-card-center_2025年08月_v2_20260811123045123456_报告_修订版.docx",
        )


if __name__ == "__main__":
    unittest.main()
