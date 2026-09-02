from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch


BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))

from app.services.audit_input import build_audit_input  # noqa: E402
from app.services.agent_worker import AgentWorker  # noqa: E402
from app.services.audit_worker import AuditWorker  # noqa: E402
from app.services.deepseek_client import (  # noqa: E402
    DeepSeekClient,
    DeepSeekNotConfigured,
    build_agent_system_content,
    build_audit_output_protocol,
    chat_completions_url,
    parse_chat_stream,
    parse_model_result,
)


def write_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("性能容量报告", level=1)
    document.add_paragraph("应用指标统计与分析内容完整。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "组件"
    table.cell(0, 1).text = "数量"
    table.cell(1, 0).text = "ES"
    table.cell(1, 1).text = "14"
    document.save(path)


class FakeAuditRepository:
    def __init__(
        self,
        job=None,
        prompts=None,
        model_config=None,
        fail_event_phase=None,
        mark_error_exception=None,
    ):
        self.job = job
        self.prompts = prompts or {}
        self.model_config = model_config or {
            "apiUrl": "https://example.test/v1",
            "modelName": "shared-model",
            "apiKey": "shared-secret",
        }
        self.fail_event_phase = fail_event_phase
        self.mark_error_exception = mark_error_exception
        self.claim_calls = 0
        self.claimed = False
        self.requested_prompt_types = []
        self.execution_contexts = []
        self.completed = []
        self.errors = []
        self.events = []
        self.jira_queues = []
        self.checkpoints = [{"id": 1, "name": "章节完整性", "content": "检查章节", "sortOrder": 10, "enabled": True}]
        self.checkpoint_reads = 0
        self.snapshots = []

    def next_pending_audit(self):
        self.claim_calls += 1
        if self.job is None or self.claimed:
            return None
        self.claimed = True
        return self.job

    def get_ai_config(self):
        return self.model_config

    def get_active_prompt(self, audit_type):
        self.requested_prompt_types.append(audit_type)
        return self.prompts.get(audit_type)

    def set_audit_execution_context(self, audit_id, prompt, model_config):
        self.execution_contexts.append(
            (audit_id, prompt["version"], model_config["modelName"])
        )

    def get_active_checkpoints(self):
        self.checkpoint_reads += 1
        return self.checkpoints

    def save_checkpoint_snapshot(self, audit_id, checkpoints):
        self.snapshots.append((audit_id, checkpoints))

    def mark_audit_complete(self, audit_id, version_id, result):
        self.completed.append((audit_id, version_id, result))

    def mark_audit_error(self, audit_id, version_id, message):
        self.errors.append((audit_id, version_id, message))
        if self.mark_error_exception:
            raise self.mark_error_exception

    def queue_jira_creation(self, report_id, version_id, audit_id):
        self.jira_queues.append((report_id, version_id, audit_id))
        return True

    def append_audit_event(self, audit_id, event_type, phase, content):
        if phase == self.fail_event_phase:
            raise RuntimeError(f"{phase} event failed")
        self.events.append((audit_id, event_type, phase, content))
        return len(self.events)


class FakeModelClient:
    def __init__(self, result=None):
        self.result = result or {
            "resultText": "审核结论：通过\nJIRA标题：容量基线审核通过\n\n## 审核总结\n未发现明显问题。",
            "conclusion": "passed",
            "jiraTitle": "容量基线审核通过",
        }
        self.calls = []

    def audit_report(self, model_config, prompt, audit_input, on_delta=None):
        self.calls.append((model_config, prompt, audit_input))
        if on_delta:
            on_delta("审核结论：通过\n")
            on_delta("未发现明显问题。")
        return self.result

    def chat_report_agent(
        self, model_config, prompt, report_context, history, question, on_delta=None
    ):
        self.calls.append((model_config, prompt, report_context, history, question))
        if on_delta:
            on_delta("建议将ES服务器数量")
            on_delta("统一为14台。")
        return "建议将ES服务器数量统一为14台。"


class FakeAgentRepository:
    def __init__(self, job=None, prompt=None, model_config=None, saved_prompt=None):
        self.job = job
        self.prompt = prompt
        self.saved_prompt = saved_prompt
        self.model_config = model_config or {
            "apiUrl": "https://example.test/v1",
            "modelName": "shared-model",
            "apiKey": "shared-secret",
        }
        self.requested_prompt_id = None
        self.requested_prompt_types = []
        self.running = []
        self.chunks = []
        self.completed = []
        self.errors = []

    def next_pending_agent_message(self):
        return self.job

    def get_ai_config(self):
        return self.model_config

    def get_prompt_by_id(self, prompt_id):
        self.requested_prompt_id = prompt_id
        return self.saved_prompt

    def get_active_prompt(self, audit_type):
        self.requested_prompt_types.append(audit_type)
        return self.prompt

    def mark_agent_message_running(self, message_id, model_name):
        self.running.append((message_id, model_name))

    def append_agent_message_content(self, message_id, content):
        self.chunks.append((message_id, content))

    def mark_agent_message_complete(self, message_id):
        self.completed.append(message_id)

    def mark_agent_message_error(self, message_id, message):
        self.errors.append((message_id, message))


class AuditWorkerTest(unittest.TestCase):
    def test_build_audit_input_from_docx(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            path = Path(temporary_dir) / "报告.docx"
            write_docx(path)
            audit_input = build_audit_input(
                {
                    "systemId": "credit-card-center",
                    "title": "报告",
                    "reportMonth": "2025年08月",
                    "fileName": "报告.docx",
                    "filePath": str(path),
                }
            )

        self.assertEqual(audit_input["report"]["reportMonth"], "2025年08月")
        self.assertFalse(audit_input["scope"]["imagesAndChartsSemantic"])
        self.assertIn("性能容量报告", audit_input["document"]["paragraphs"])
        self.assertEqual(audit_input["document"]["tables"][0][1], ["ES", "14"])
        self.assertEqual(audit_input["document"]["headings"][0]["level"], 1)
        self.assertEqual(audit_input["document"]["headings"][0]["text"], "性能容量报告")
        self.assertFalse(audit_input["document"]["tocAvailable"])
        self.assertEqual(audit_input["document"]["tocEntries"], [])
        self.assertNotIn("checkpoints", audit_input)

    def test_parse_model_result_extracts_optional_conclusion(self):
        self.assertEqual(parse_model_result("审核结论：通过\n内容完整")["conclusion"], "passed")
        self.assertEqual(parse_model_result("审核结论：不通过\n存在问题")["conclusion"], "failed")
        self.assertEqual(parse_model_result("审核工作已经完成")["conclusion"], "completed")

    def test_parse_model_result_uses_last_standalone_conclusion(self):
        result = parse_model_result(
            "历史说明中提到审核结论：不通过，但本次未发现问题。\n"
            "审核结论：不通过\n"
            "复核后确认无实际问题。\n"
            "审核结论：通过"
        )

        self.assertEqual(result["conclusion"], "passed")

    def test_parse_model_result_extracts_and_limits_jira_title(self):
        result = parse_model_result(
            "审核结论：通过\nJIRA标题：信用卡中心性能容量报告审核通过\n内容完整"
        )

        self.assertEqual(result["jiraTitle"], "信用卡中心性能容量报告审核通过")
        self.assertLessEqual(len(result["jiraTitle"]), 15)

    def test_parse_model_result_extracts_jira_title_after_long_markdown(self):
        result = parse_model_result(
            "## 审核总结\n" + "报告内容正常。" * 100 + "\n"
            "JIRA标题：容量报告初审通过\n"
            "审核结论：通过"
        )

        self.assertEqual(result["jiraTitle"], "容量报告初审通过")
        self.assertEqual(result["conclusion"], "passed")

    def test_audit_protocol_requires_evidence_consistent_conclusion(self):
        protocol = build_audit_output_protocol("initial")

        self.assertIn("未发现确认问题，必须判定为通过", protocol)
        self.assertIn("无法判断、证据不足", protocol)
        self.assertIn("不通过时必须列出至少一项确认问题", protocol)
        self.assertIn("JIRA标题", protocol)

    def test_agent_treats_latest_audit_result_as_reviewable_history(self):
        system_content = build_agent_system_content(
            {"promptContent": "检查报告"},
            {"latestAuditResult": "审核结论：不通过", "document": {}},
        )

        self.assertIn("只是系统已保存的历史审核记录", system_content)
        self.assertIn("独立核对当前报告", system_content)
        self.assertIn("系统已保存结论", system_content)
        self.assertIn("通过重新审核产生正式新结论", system_content)

    def test_chat_completions_url_accepts_base_or_full_url(self):
        self.assertEqual(
            chat_completions_url("https://ark.cn-beijing.volces.com/api/coding/v3"),
            "https://ark.cn-beijing.volces.com/api/coding/v3/chat/completions",
        )

    def test_parse_chat_stream_collects_content_deltas(self):
        deltas = []
        lines = [
            b'data: {"choices":[{"delta":{"role":"assistant"}}]}\n',
            b'data: {"choices":[{"delta":{"content":"{\\"summary\\":"}}]}\n',
            b'data: {"choices":[{"delta":{"content":"{}}"}}]}\n',
            b'data: {"choices":[],"usage":{"total_tokens":1024}}\n',
            b'data: [DONE]\n',
        ]

        content = parse_chat_stream(lines, deltas.append)

        self.assertEqual(content, '{"summary":{}}')
        self.assertEqual(deltas, ['{"summary":', '{}}'])
        self.assertEqual(
            chat_completions_url("https://ark.cn-beijing.volces.com/api/coding/v3/chat/completions"),
            "https://ark.cn-beijing.volces.com/api/coding/v3/chat/completions",
        )

    def test_worker_completes_pending_audit(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            path = Path(temporary_dir) / "报告.docx"
            write_docx(path)
            job = {
                "auditId": 99,
                "reportId": 1,
                "versionId": 10,
                "filePath": str(path),
                "fileName": "报告.docx",
                "systemId": "credit-card-center",
                "title": "报告",
                "reportMonth": "2025年08月",
                "auditType": "initial",
            }
            prompt = {"id": 1, "version": 2, "promptContent": "请审核"}
            repository = FakeAuditRepository(job, {"initial": prompt})
            model_client = FakeModelClient()

            with self.assertLogs("app.services.audit_worker", level="INFO") as captured:
                result = AuditWorker(repository, model_client).run_once()

        self.assertEqual(result, {"processed": True, "auditId": 99, "status": "completed"})
        self.assertEqual(repository.execution_contexts, [(99, 2, "shared-model")])
        self.assertEqual(repository.completed[0][0:2], (99, 10))
        self.assertEqual(repository.requested_prompt_types, ["initial"])
        self.assertEqual(repository.checkpoint_reads, 0)
        self.assertEqual(repository.snapshots, [])
        self.assertEqual(model_client.calls[0][0]["modelName"], "shared-model")
        self.assertEqual(model_client.calls[0][2]["report"]["systemId"], "credit-card-center")
        self.assertNotIn("checkpoints", model_client.calls[0][2])
        self.assertEqual(repository.events[0][2], "started")
        self.assertTrue(any(event[1] == "model" for event in repository.events))
        self.assertEqual(repository.events[-1][2], "completed")
        self.assertEqual(repository.jira_queues, [(1, 10, 99)])
        audit_logs = "\n".join(captured.output)
        self.assertIn("audit_completed reportId=1 versionId=10 auditId=99", audit_logs)
        self.assertIn("jira_queue_requested reportId=1 versionId=10 auditId=99", audit_logs)
        self.assertNotIn("shared-secret", audit_logs)

    def test_initial_pass_without_jira_title_still_queues_jira(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            path = Path(temporary_dir) / "报告.docx"
            write_docx(path)
            repository = FakeAuditRepository(
                {
                    "auditId": 99,
                    "reportId": 1,
                    "versionId": 10,
                    "auditType": "initial",
                    "filePath": str(path),
                    "fileName": "报告.docx",
                    "systemId": "credit-card-center",
                    "title": "报告",
                    "reportMonth": "2026年08月",
                },
                {"initial": {"id": 1, "version": 1, "promptContent": "请审核"}},
            )
            model_client = FakeModelClient(
                {
                    "resultText": "审核结论：通过\n内容完整",
                    "conclusion": "passed",
                    "jiraTitle": "",
                }
            )

            result = AuditWorker(repository, model_client).run_once()

        self.assertEqual(result["status"], "completed")
        self.assertEqual(repository.errors, [])
        self.assertEqual(repository.jira_queues, [(1, 10, 99)])

    def test_worker_handles_no_pending_audit(self):
        result = AuditWorker(FakeAuditRepository(), FakeModelClient()).run_once()

        self.assertEqual(result, {"processed": False, "reason": "no_pending_audit"})

    def test_worker_marks_error_when_type_prompt_missing(self):
        for audit_type, label in (("initial", "初始"), ("revision", "修订")):
            with self.subTest(audit_type=audit_type):
                repository = FakeAuditRepository(
                    {"auditId": 99, "versionId": 10, "auditType": audit_type}
                )

                result = AuditWorker(repository, FakeModelClient()).run_once()

                message = f"未配置启用的{label}审核提示词"
                self.assertEqual(result, {"processed": True, "auditId": 99, "status": "error"})
                self.assertEqual(repository.errors, [(99, 10, message)])
                self.assertEqual(repository.requested_prompt_types, [audit_type])
                self.assertEqual(repository.events[-1][2], "configuration")

    def test_worker_marks_error_when_shared_model_config_missing_without_leaking_key(self):
        repository = FakeAuditRepository(
            {"auditId": 99, "versionId": 10, "auditType": "initial"},
            {"initial": {"id": 1, "version": 1, "promptContent": "请审核"}},
            model_config={},
        )
        repository.model_config = {"apiUrl": "", "modelName": "", "apiKey": "secret-must-not-leak"}

        result = AuditWorker(repository, FakeModelClient()).run_once()

        self.assertEqual(result, {"processed": True, "auditId": 99, "status": "error"})
        self.assertEqual(repository.errors, [(99, 10, "未配置共用模型连接")])
        self.assertNotIn("secret-must-not-leak", str(repository.events))
        self.assertEqual(repository.requested_prompt_types, [])

    def test_worker_marks_error_when_file_missing(self):
        job = {
            "auditId": 99,
            "versionId": 10,
            "filePath": "/not/exist.docx",
            "fileName": "报告.docx",
            "systemId": "credit-card-center",
            "title": "报告",
            "reportMonth": "2025年08月",
            "auditType": "revision",
        }
        prompt = {"id": 1, "version": 2, "promptContent": "请审核"}
        repository = FakeAuditRepository(job, {"revision": prompt})

        result = AuditWorker(repository, FakeModelClient()).run_once()

        self.assertEqual(result, {"processed": True, "auditId": 99, "status": "error"})
        self.assertEqual(repository.errors, [(99, 10, "报告文件不存在")])
        self.assertEqual(repository.events[-1][2], "failed")

    def test_worker_selects_only_the_revision_prompt(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            path = Path(temporary_dir) / "报告.docx"
            write_docx(path)
            job = {
                "auditId": 100,
                "reportId": 1,
                "versionId": 11,
                "filePath": str(path),
                "fileName": "报告.docx",
                "systemId": "credit-card-center",
                "title": "报告",
                "reportMonth": "2025年08月",
                "auditType": "revision",
            }
            prompts = {
                "initial": {"id": 1, "version": 1, "promptContent": "初始提示词"},
                "revision": {"id": 2, "version": 3, "promptContent": "修订提示词"},
            }
            repository = FakeAuditRepository(job, prompts)
            model_client = FakeModelClient()

            result = AuditWorker(repository, model_client).run_once()

        self.assertEqual(result["status"], "completed")
        self.assertEqual(repository.requested_prompt_types, ["revision"])
        self.assertEqual(model_client.calls[0][1]["promptContent"], "修订提示词")
        self.assertNotIn("checkpoints", model_client.calls[0][2])
        self.assertEqual(repository.jira_queues, [])

    def test_worker_executes_a_claimed_job_only_once(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            path = Path(temporary_dir) / "报告.docx"
            write_docx(path)
            job = {
                "auditId": 99,
                "reportId": 1,
                "versionId": 10,
                "filePath": str(path),
                "fileName": "报告.docx",
                "systemId": "credit-card-center",
                "title": "报告",
                "reportMonth": "2025年08月",
                "auditType": "initial",
            }
            repository = FakeAuditRepository(
                job,
                {"initial": {"id": 1, "version": 2, "promptContent": "请审核"}},
            )
            model_client = FakeModelClient()
            worker = AuditWorker(repository, model_client)

            first = worker.run_once()
            second = worker.run_once()

        self.assertEqual(first["status"], "completed")
        self.assertEqual(second, {"processed": False, "reason": "no_pending_audit"})
        self.assertEqual(repository.claim_calls, 2)
        self.assertEqual(len(model_client.calls), 1)
        self.assertEqual(len(repository.completed), 1)

    def test_worker_marks_error_when_started_event_fails(self):
        repository = FakeAuditRepository(
            {"auditId": 99, "versionId": 10, "auditType": "initial"},
            {"initial": {"id": 1, "version": 2, "promptContent": "请审核"}},
            fail_event_phase="started",
        )

        result = AuditWorker(repository, FakeModelClient()).run_once()

        self.assertEqual(result, {"processed": True, "auditId": 99, "status": "error"})
        self.assertEqual(repository.errors, [(99, 10, "started event failed")])

    def test_worker_marks_error_when_prompt_fields_are_invalid(self):
        repository = FakeAuditRepository(
            {"auditId": 99, "versionId": 10, "auditType": "initial"},
            {"initial": {"id": 1, "promptContent": "请审核"}},
        )

        result = AuditWorker(repository, FakeModelClient()).run_once()

        self.assertEqual(result, {"processed": True, "auditId": 99, "status": "error"})
        self.assertEqual(repository.errors, [(99, 10, "'version'")])

    def test_worker_returns_stable_error_when_error_persistence_also_fails(self):
        repository = FakeAuditRepository(
            {"auditId": 99, "versionId": 10, "auditType": "initial"},
            {"initial": {"id": 1, "version": 2, "promptContent": "请审核"}},
            fail_event_phase="started",
            mark_error_exception=RuntimeError("error persistence failed"),
        )

        with self.assertLogs("app.services.audit_worker", level="ERROR") as logs:
            result = AuditWorker(repository, FakeModelClient()).run_once()

        self.assertEqual(result, {"processed": True, "auditId": 99, "status": "error"})
        self.assertTrue(any("started event failed" in line for line in logs.output))
        self.assertEqual(repository.errors, [(99, 10, "started event failed")])

    def test_audit_report_uses_shared_model_config_and_runtime_output_protocol(self):
        settings = SimpleNamespace(
            deepseek_url="https://legacy-env.test/v1",
            deepseek_key="legacy-env-secret",
            deepseek_model="legacy-env-model",
        )
        client = DeepSeekClient(settings)
        response_lines = [
            (
                "data: "
                + json.dumps(
                    {"choices": [{"delta": {"content": "审核结论：通过\n\n## 审核总结\n内容完整。"}}]},
                    ensure_ascii=False,
                )
                + "\n"
            ).encode("utf-8"),
            b"data: [DONE]\n",
        ]

        class Response:
            def __enter__(self):
                return iter(response_lines)

            def __exit__(self, exc_type, exc_value, traceback):
                return False

        with patch("app.services.deepseek_client.urllib.request.urlopen", return_value=Response()) as urlopen:
            result = client.audit_report(
                {
                    "apiUrl": "https://shared-config.test/v1",
                    "apiKey": "shared-config-secret",
                    "modelName": "shared-config-model",
                },
                {
                    "promptContent": "只审核可读取内容",
                    "auditType": "initial",
                    "apiUrl": "https://legacy-prompt.test/v1",
                    "apiKey": "legacy-prompt-secret",
                    "modelName": "legacy-prompt-model",
                },
                {"report": {"title": "报告"}},
            )

        request = urlopen.call_args.args[0]
        payload = json.loads(request.data.decode("utf-8"))
        system_message = payload["messages"][0]["content"]
        user_message = payload["messages"][1]["content"]
        self.assertEqual(request.full_url, "https://shared-config.test/v1/chat/completions")
        self.assertEqual(request.get_header("Authorization"), "Bearer shared-config-secret")
        self.assertEqual(payload["model"], "shared-config-model")
        self.assertIn("最后一行必须且只能是“审核结论：通过”或“审核结论：不通过”", system_message)
        self.assertIn("JIRA标题", system_message)
        self.assertIn("不超过15个中文字符", system_message)
        self.assertIn("先使用中文 Markdown", system_message)
        self.assertIn("不要输出 JSON", system_message)
        self.assertNotIn("检查点", user_message)
        self.assertEqual(result["conclusion"], "passed")

    def test_audit_report_does_not_fall_back_to_prompt_or_environment_model_config(self):
        settings = SimpleNamespace(
            deepseek_url="https://legacy-env.test/v1",
            deepseek_key="legacy-env-secret",
            deepseek_model="legacy-env-model",
        )
        client = DeepSeekClient(settings)

        with patch("app.services.deepseek_client.urllib.request.urlopen") as urlopen:
            with self.assertRaisesRegex(DeepSeekNotConfigured, "共用模型连接") as raised:
                client.audit_report(
                    {"apiUrl": "", "apiKey": "", "modelName": ""},
                    {
                        "promptContent": "审核提示词",
                        "apiUrl": "https://legacy-prompt.test/v1",
                        "apiKey": "legacy-prompt-secret",
                        "modelName": "legacy-prompt-model",
                    },
                    {"report": {"title": "报告"}},
                )

        urlopen.assert_not_called()
        self.assertNotIn("legacy-prompt-secret", str(raised.exception))
        self.assertNotIn("legacy-env-secret", str(raised.exception))

    def test_agent_worker_streams_and_completes_reply(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            path = Path(temporary_dir) / "报告.docx"
            write_docx(path)
            job = {
                "messageId": 21,
                "filePath": str(path),
                "fileName": "报告.docx",
                "systemId": "capacity",
                "title": "容量报告",
                "reportMonth": "2026年08月",
                "question": "ES服务器数量应该怎么改？",
                "auditResult": "审核结论：不通过",
                "auditType": "revision",
                "promptId": 42,
                "history": [{"role": "user", "content": "报告有什么问题？"}],
            }
            prompt = {"id": 42, "auditType": "revision", "promptContent": "请审核"}
            repository = FakeAgentRepository(job, saved_prompt=prompt)
            model_client = FakeModelClient()

            result = AgentWorker(repository, model_client).run_once()

        self.assertEqual(result, {"processed": True, "messageId": 21, "status": "completed"})
        self.assertEqual(repository.running, [(21, "shared-model")])
        self.assertEqual("".join(chunk for _, chunk in repository.chunks), "建议将ES服务器数量统一为14台。")
        self.assertEqual(repository.completed, [21])
        self.assertEqual(repository.requested_prompt_id, 42)
        self.assertEqual(repository.requested_prompt_types, [])
        self.assertEqual(model_client.calls[0][0]["modelName"], "shared-model")
        self.assertEqual(model_client.calls[0][1]["auditType"], "revision")
        self.assertEqual(model_client.calls[0][4], "ES服务器数量应该怎么改？")
        self.assertNotIn("checkpoints", model_client.calls[0][2])

    def test_agent_worker_marks_error_when_model_is_missing(self):
        repository = FakeAgentRepository(
            {"messageId": 22, "auditType": "revision", "promptId": None},
            model_config={},
        )
        repository.model_config = {}

        result = AgentWorker(repository, FakeModelClient()).run_once()

        self.assertEqual(result, {"processed": True, "messageId": 22, "status": "error"})
        self.assertEqual(repository.errors, [(22, "未配置共用模型连接")])


if __name__ == "__main__":
    unittest.main()
